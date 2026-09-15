"""Tests for plugins/user/excel_tool.py — Excel tools (create/inspect/read/modify)."""

import json
import os
from datetime import datetime

import openpyxl
import pytest

from plugins.user.excel_tool import (
    ExcelCreateTool,
    ExcelDocument,
    ExcelInspectTool,
    ExcelModifyTool,
    ExcelReadTool,
)

SALES_SHEET = {
    "name": "Sales",
    "columns": [
        {"name": "日期", "type": "date"},
        {"name": "产品", "type": "text"},
        {"name": "数量", "type": "integer"},
        {"name": "金额", "type": "currency"},
    ],
    "rows": [
        ["2026-09-01", "A", 100, 12000],
        ["2026-09-02", "B", 200, 18000],
    ],
}

SUMMARY_SHEET = {
    "name": "Summary",
    "columns": [{"name": "指标", "type": "text"}, {"name": "数值", "type": "float"}],
    "rows": [["总金额", 30000.0]],
}

TWO_SHEETS = [SALES_SHEET, SUMMARY_SHEET]


@pytest.fixture
def xlsx_path(tmp_path):
    """在当前临时目录创建一份基础文件，返回其绝对路径。"""
    file = str(tmp_path / "sales_report.xlsx")
    result = json.loads(ExcelCreateTool().execute(filename=file, sheets=TWO_SHEETS))
    assert result["success"] is True
    return file


def _load(file):
    return openpyxl.load_workbook(file, data_only=False)


# ── excel_create ────────────────────────────────────────────────────────────


class TestExcelCreate:
    def test_create_returns_path_and_writes_file(self, tmp_path):
        file = str(tmp_path / "a.xlsx")
        result = json.loads(ExcelCreateTool().execute(filename=file, sheets=TWO_SHEETS))
        assert result["success"] is True
        assert result["path"] == file
        assert result["sheets_created"] == 2
        assert os.path.exists(file)

        wb = _load(file)
        assert wb.sheetnames == ["Sales", "Summary"]
        ws = wb["Sales"]
        assert ws.cell(1, 1).value == "日期"
        assert ws.cell(1, 4).value == "金额"
        assert ws.freeze_panes == "A2"

    def test_create_type_coercion_and_number_format(self, tmp_path):
        file = str(tmp_path / "types.xlsx")
        sheets = [{
            "name": "T",
            "columns": [
                {"name": "日期", "type": "date"},
                {"name": "百分比", "type": "percent"},
                {"name": "布尔", "type": "boolean"},
                {"name": "金额", "type": "currency"},
            ],
            "rows": [
                ["2026-09-01", "12.5", "是", 1234.5],
            ],
        }]
        result = json.loads(ExcelCreateTool().execute(filename=file, sheets=sheets))
        assert result["success"] is True

        ws = _load(file)["T"]
        assert isinstance(ws.cell(2, 1).value, datetime)
        assert ws.cell(2, 1).number_format == "yyyy-mm-dd"
        assert ws.cell(2, 2).value == pytest.approx(0.125)
        assert ws.cell(2, 2).number_format == "0.0%"
        assert ws.cell(2, 3).value is True
        assert ws.cell(2, 4).number_format == "#,##0.00"

    def test_row_column_mismatch_reports_location(self, tmp_path):
        file = str(tmp_path / "bad.xlsx")
        sheets = [{
            "name": "S",
            "columns": [{"name": "A", "type": "text"}, {"name": "B", "type": "text"}],
            "rows": [["只有一列"]],
        }]
        result = json.loads(ExcelCreateTool().execute(filename=file, sheets=sheets))
        assert result["success"] is False
        assert "列数" in result["error"]

    def test_overwrite_protection(self, tmp_path):
        file = str(tmp_path / "o.xlsx")
        ExcelCreateTool().execute(filename=file, sheets=TWO_SHEETS)
        again = json.loads(ExcelCreateTool().execute(filename=file, sheets=TWO_SHEETS))
        assert again["success"] is False
        assert "overwrite" in again["error"]

        forced = json.loads(
            ExcelCreateTool().execute(filename=file, sheets=TWO_SHEETS, overwrite=True)
        )
        assert forced["success"] is True

    def test_missing_extension_appended(self, tmp_path):
        file = str(tmp_path / "noext")
        result = json.loads(ExcelCreateTool().execute(filename=file, sheets=TWO_SHEETS))
        assert result["success"] is True
        assert result["path"].endswith(".xlsx")
        assert os.path.exists(file + ".xlsx")


# ── excel_inspect ───────────────────────────────────────────────────────────


class TestExcelInspect:
    def test_inspect_structure(self, xlsx_path):
        result = json.loads(ExcelInspectTool().execute(file=xlsx_path))
        assert result["success"] is True
        assert result["sheets"] == [
            {"name": "Sales", "rows": 3, "columns": 4,
             "headers": ["日期", "产品", "数量", "金额"]},
            {"name": "Summary", "rows": 2, "columns": 2,
             "headers": ["指标", "数值"]},
        ]

    def test_inspect_missing_file(self, tmp_path):
        result = json.loads(
            ExcelInspectTool().execute(file=str(tmp_path / "nope.xlsx"))
        )
        assert result["success"] is False
        assert "文件不存在" in result["error"]

    def test_inspect_without_extension_resolves(self, xlsx_path):
        bare = xlsx_path[:-5]  # 去掉 .xlsx
        result = json.loads(ExcelInspectTool().execute(file=bare))
        assert result["success"] is True
        assert result["path"] == xlsx_path


# ── excel_read ──────────────────────────────────────────────────────────────


class TestExcelRead:
    def test_read_returns_dict_rows(self, xlsx_path):
        result = json.loads(ExcelReadTool().execute(file=xlsx_path))
        assert result["success"] is True
        assert result["sheet"] == "Sales"
        assert result["row_count"] == 2
        assert result["returned"] == 2
        assert result["start_row"] == 1
        assert result["data"] == [
            {"日期": "2026-09-01", "产品": "A", "数量": 100, "金额": 12000},
            {"日期": "2026-09-02", "产品": "B", "数量": 200, "金额": 18000},
        ]

    def test_read_selected_columns(self, xlsx_path):
        result = json.loads(
            ExcelReadTool().execute(
                file=xlsx_path, sheet="Sales", columns=["产品", "金额"], limit=5
            )
        )
        assert result["success"] is True
        assert result["columns"] == ["产品", "金额"]
        assert result["data"][0] == {"产品": "A", "金额": 12000}

    def test_read_pagination_offset_and_limit(self, xlsx_path):
        result = json.loads(
            ExcelReadTool().execute(file=xlsx_path, sheet="Sales", limit=1, offset=1)
        )
        assert result["success"] is True
        assert result["returned"] == 1
        assert result["start_row"] == 2
        assert result["data"] == [{"日期": "2026-09-02", "产品": "B", "数量": 200, "金额": 18000}]

    def test_read_unknown_column_reports_headers(self, xlsx_path):
        result = json.loads(
            ExcelReadTool().execute(file=xlsx_path, sheet="Sales", columns=["不存在的列"])
        )
        assert result["success"] is False
        assert "可用表头" in result["error"]

    def test_read_sheet_names_available(self, xlsx_path):
        result = json.loads(
            ExcelReadTool().execute(file=xlsx_path, sheet="Summary")
        )
        assert result["success"] is True
        assert result["data"] == [{"指标": "总金额", "数值": 30000.0}]


# ── excel_modify ────────────────────────────────────────────────────────────


class TestExcelModify:
    def test_append_rows_lists(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[{
                    "op": "append_rows",
                    "sheet": "Sales",
                    "rows": [["2026-09-15", "A", 100, 2000]],
                }],
            )
        )
        assert result["success"] is True
        op = result["operations"][0]
        assert op["success"] is True
        assert op["rows_appended"] == 1
        assert op["start_row"] == 4

        ws = _load(xlsx_path)["Sales"]
        assert ws.cell(4, 1).value == datetime(2026, 9, 15)
        assert ws.cell(4, 4).value == 2000

    def test_append_rows_dicts(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[{
                    "op": "append_rows",
                    "sheet": "Sales",
                    "rows": [{"日期": "2026-09-16", "产品": "C", "数量": 50, "金额": 999}],
                }],
            )
        )
        assert result["success"] is True
        assert result["operations"][0]["start_row"] == 4
        ws = _load(xlsx_path)["Sales"]
        assert list(ws.cell(4, c).value for c in range(1, 5)) == [
            datetime(2026, 9, 16), "C", 50, 999,
        ]

    def test_add_formula_bare_names_translated(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[{
                    "op": "add_formula",
                    "sheet": "Sales",
                    "column": "利润",
                    "formula": "=数量*金额",
                }],
            )
        )
        assert result["success"] is True
        op = result["operations"][0]
        assert op["success"] is True
        assert op["rows_affected"] == 2

        ws = _load(xlsx_path)["Sales"]
        assert ws.cell(1, 5).value == "利润"
        assert ws.cell(2, 5).value == "=C2*D2"
        assert ws.cell(3, 5).value == "=C3*D3"

    def test_add_formula_placeholder_and_standard_ref_pass_through(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[{
                    "op": "add_formula",
                    "sheet": "Sales",
                    "column": "小计",
                    "formula": "={数量}*$D$1+SUM(B2:B3)",
                }],
            )
        )
        assert result["success"] is True
        ws = _load(xlsx_path)["Sales"]
        assert ws.cell(2, 5).value == "=C2*$D$1+SUM(B2:B3)"

    def test_add_formula_range_limits_rows(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[{
                    "op": "add_formula",
                    "sheet": "Sales",
                    "column": "利润",
                    "formula": "=数量*金额",
                    "rows": "2:2",
                }],
            )
        )
        assert result["success"] is True
        assert result["operations"][0]["rows_affected"] == 1
        ws = _load(xlsx_path)["Sales"]
        assert ws.cell(2, 5).value == "=C2*D2"
        assert ws.cell(3, 5).value is None

    def test_add_chart_with_auto_created_sheet(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[{
                    "op": "add_chart",
                    "sheet": "图表",
                    "type": "bar",
                    "title": "产品销售额",
                    "category": "产品",
                    "value": "金额",
                }],
            )
        )
        assert result["success"] is True
        op = result["operations"][0]
        assert op["success"] is True
        assert op["sheet_created"] is True

        wb = _load(xlsx_path)
        assert "图表" in wb.sheetnames
        charts = getattr(wb["图表"], "_charts")
        assert len(charts) == 1
        title_text = charts[0].title.tx.rich.p[0].r[0].t
        assert title_text == "产品销售额"

    def test_partial_op_failure_still_saves_others(self, xlsx_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=xlsx_path,
                operations=[
                    {"op": "append_rows", "sheet": "Sales",
                     "rows": [["2026-09-15", "A", 100, 2000]]},
                    {"op": "append_rows", "sheet": "不存在的表", "rows": [["x"]]},
                    {"op": "add_formula", "sheet": "Sales",
                     "column": "利润", "formula": "=金额/100"},
                ],
            )
        )
        assert result["success"] is True
        ops = result["operations"]
        assert ops[0]["success"] is True
        assert ops[1]["success"] is False
        assert "不存在" in ops[1]["error"]
        assert ops[2]["success"] is True

        ws = _load(xlsx_path)["Sales"]
        assert ws.cell(4, 1).value == datetime(2026, 9, 15)
        assert ws.cell(1, 5).value == "利润"
        assert ws.cell(2, 5).value == "=D2/100"

    def test_modify_missing_file(self, tmp_path):
        result = json.loads(
            ExcelModifyTool().execute(
                file=str(tmp_path / "ghost.xlsx"),
                operations=[{"op": "append_rows", "sheet": "S", "rows": [["x"]]}],
            )
        )
        assert result["success"] is False
        assert "文件不存在" in result["error"]


# ── 路径解析 ────────────────────────────────────────────────────────────────


class TestExcelDocumentPath:
    def test_resolve_absolute_unchanged(self):
        assert ExcelDocument.resolve_path("/tmp/abc.xlsx") == "/tmp/abc.xlsx"

    def test_resolve_bare_filename_to_output_dir(self):
        path = ExcelDocument.resolve_path("report.xlsx")
        assert path.endswith(os.path.join("output", "report.xlsx"))
        assert os.path.isabs(path)

    def test_resolve_relative_with_dir_anchored_to_project(self):
        from modules.utils.paths import PROJECT_ROOT
        path = ExcelDocument.resolve_path("sub/dir.xlsx")
        assert path == os.path.join(PROJECT_ROOT, "sub", "dir.xlsx")

    def test_resolve_empty_raises(self):
        with pytest.raises(ValueError):
            ExcelDocument.resolve_path("")